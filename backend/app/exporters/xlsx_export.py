"""Экспорт таблицы в XLSX через openpyxl."""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.exporters.table import ExportTable


def render_xlsx(table: ExportTable, *, sheet_name: str = "Данные") -> bytes:
    """Книга Excel: заголовок жирным, автоширина, типы чисел сохраняются."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = (sheet_name or "Данные")[:31]

    headers = table.header()
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for row in table.as_matrix():
        sheet.append([_excel_value(value) for value in row])

    for index, key in enumerate(table.columns, start=1):
        width = max(len(str(headers[index - 1])), 8)
        for row in table.rows[:200]:
            width = max(width, min(len(str(row.get(key, "") or "")), 48))
        sheet.column_dimensions[get_column_letter(index)].width = width + 2

    sheet.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _excel_value(value: Any) -> Any:
    if value == "":
        return None
    return value
